#!/usr/bin/env python3
"""Cross-artifact alignment audit for the C398 courseware package.

Verifies that the slide deck, Lesson Plan, Learner Guide (DOCX + Markdown) and
the labs/ folder all agree - lab numbering, step counts, topics, learning
outcomes, schedule minutes and the non-WSQ structure. Exit 0 = aligned.

Run:  python tools/audit_alignment.py
"""
import re, sys, glob, os, importlib
from pptx import Presentation
from docx import Document

HERE=os.path.dirname(os.path.abspath(__file__))
REPO=os.path.dirname(HERE)
BUILD=os.path.join(REPO,".claude","skills","non-wsq-courseware-build","build")
sys.path.insert(0,BUILD)
import course_data as C

ACT=[]
for f in sorted(glob.glob(os.path.join(BUILD,"data_domain[0-9]*.py")),
                key=lambda q:int("".join(c for c in os.path.basename(q) if c.isdigit()) or 0)):
    n="".join(c for c in os.path.basename(f) if c.isdigit())
    ACT+=getattr(importlib.import_module(os.path.basename(f)[:-3]),f"DOMAIN{n}")

CW=os.path.join(REPO,"courseware")
PPTX=os.path.join(CW,f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
LPX =os.path.join(CW,f"LP-{C.SHORT_TITLE}.docx")
LGX =os.path.join(CW,f"LG-{C.SHORT_TITLE}.docx")
LGMD=os.path.join(REPO,f"LG-{C.SHORT_TITLE}.md")

fails=[]; oks=[]
def chk(cond,msg): (oks if cond else fails).append(msg)

# ---------------------------------------------------------------- source
chk(all(a["num"]==i for i,a in enumerate(ACT,1)), f"lab numbering contiguous 1..{len(ACT)}")
chk(len({a["title"] for a in ACT})==len(ACT), "lab titles unique")
for a in ACT:
    chk(len(a.get("step_details",[]))==len(a["steps"]),
        f"Lab {a['num']}: step_details count == steps count")

# ---------------------------------------------------------------- deck
prs=Presentation(PPTX)
ppt=["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
allppt="\n".join(ppt)
chk(C.COURSE_CODE in ppt[0] and C.TITLE in ppt[0], "deck cover carries title + course code")
chk(C.VERSION in ppt[0], f"deck cover carries version {C.VERSION}")
for a in ACT:
    chk(any(f"LAB {a['num']}" in t and a["title"] in t for t in ppt), f"deck has Lab {a['num']} slides")
    steps=[t for t in ppt if f"LAB {a['num']}" in t and "STEP " in t]
    chk(len(steps)==len(a["steps"]),
        f"deck Lab {a['num']}: {len(steps)} step slides == {len(a['steps'])} source steps")
    chk(any(f"LAB {a['num']} · VERIFY" in t for t in ppt), f"deck Lab {a['num']} has a Test it slide")
chk("How You'll Learn" in allppt, "deck has 'How You'll Learn' (non-WSQ replacement for the assessment block)")
chk("Briefing for Assessment" not in allppt and "Assessment Flow" not in allppt, "deck has no assessment admin slides")
chk("Digital Attendance" not in allppt and "TRAQOM" not in allppt, "deck has no attendance/TRAQOM slides")
chk("Thank You!" in ppt[-1], "deck closes on Thank You")
for t in C.TOPICS:
    chk(any(f"Recap — {t['title']}" in x for x in ppt), f"deck has a recap for Topic {t['code']}")

# ---------------------------------------------------------------- lesson plan
lp=Document(LPX)
lptext="\n".join(p.text for p in lp.paragraphs)
lptab="\n".join(c.text for t in lp.tables for r in t.rows for c in r.cells)
chk(C.COURSE_CODE in lptab, "LP carries the course code")
chk(f"{C.DAYS} day" in lptab and f"{C.HOURS_PER_DAY:g} instructional hours" in lptab,
    f"LP duration = {C.DAYS} day / {C.HOURS_PER_DAY:g} hours")
for lo in C.LEARNING_OUTCOMES:
    chk(lo in lptext, f"LP carries {lo.split(':')[0]}")
sched=C.SCHEDULE(lambda n:"")
for day,(theme,rows) in sched.items():
    total=sum(m for _s,_e,m,k,_t in rows if k!="lunch")
    tea=sum(m for _s,_e,m,k,_t in rows if k=="break")
    chk(total==480, f"LP day {day} totals 480 min excluding lunch (got {total})")
    chk(total-tea==C.HOURS_PER_DAY*60, f"LP day {day} instructional = {C.HOURS_PER_DAY*60:g} min (got {total-tea})")
    chk(not any(k=="assess" for _s,_e,_m,k,_t in rows), f"LP day {day} has no assessment block")
for a in ACT:
    chk(f"Lab {a['num']}: {a['title']}" in lptab, f"LP schedule names Lab {a['num']}")
chk("Learning Reinforcement" in lptext, "LP has Learning Reinforcement (not assessment)")

# ---------------------------------------------------------------- learner guide
lg=Document(LGX)
lgtext="\n".join(p.text for p in lg.paragraphs)
md=open(LGMD,encoding="utf-8").read()
for a in ACT:
    chk(f"Lab {a['num']} — {a['title']}" in lgtext, f"LG DOCX has the Lab {a['num']} section")
    chk(f"Lab {a['num']} — {a['title']}" in md, f"LG Markdown has the Lab {a['num']} section")
    chk(a["test"][:60] in lgtext, f"LG Lab {a['num']} carries its Test it text")
for t in C.TOPICS:
    chk(f"Topic {t['code']} — {t['title']}" in lgtext, f"LG has Topic {t['code']}")
    for c_ in t["concepts"]:
        chk(c_ in lgtext, f"LG carries a Topic {t['code']} key concept")
chk(all(lo in lgtext for lo in C.LEARNING_OUTCOMES), "LG carries every learning outcome")
chk(bool(C.LG_GLOSSARY) and C.LG_GLOSSARY[0][0] in lgtext, "LG has the glossary")

# ---------------------------------------------------------------- labs
for a in ACT:
    d=glob.glob(os.path.join(REPO,"labs",f"lab{a['num']:02d}-*","README.md"))
    chk(len(d)==1, f"labs/ has exactly one folder for Lab {a['num']}")
    if d:
        txt=open(d[0],encoding="utf-8").read()
        for sec in ["## Goal","## What you'll build","## Prerequisites","## Steps","## Test it",
                    "## Troubleshooting","## Challenge","## Reflection"]:
            chk(sec in txt, f"Lab {a['num']} README has {sec}")
        chk(C.COURSE_CODE in txt, f"Lab {a['num']} README footer carries {C.COURSE_CODE}")
        chk("TGS-" not in txt, f"Lab {a['num']} README carries no TGS- reference")
        chk(len(re.findall(r"^\d+\. \*\*", txt, re.M))==len(a["steps"]),
            f"Lab {a['num']} README step count == source ({len(a['steps'])})")
idx=open(os.path.join(REPO,"labs","README.md"),encoding="utf-8").read()
for a in ACT: chk(a["title"] in idx, f"labs/README.md indexes Lab {a['num']}")
lab_block=sum(m for _s,_e,m,k,_t in sched[1][1] if k=="lab")
chk(sum(a["duration"] for a in ACT)==lab_block,
    f"lab durations ({sum(a['duration'] for a in ACT)}) == LP lab blocks ({lab_block})")

print("\n".join("  OK  "+m for m in oks))
if fails:
    print("\nFAILURES:"); print("\n".join("  XX  "+m for m in fails))
print(f"\n{len(oks)} passed, {len(fails)} failed")
sys.exit(1 if fails else 0)
